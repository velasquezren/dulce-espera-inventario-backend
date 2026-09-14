import hashlib
import html
import time
import uuid
from datetime import datetime
from typing import List
from fastapi import FastAPI, Depends, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from starlette.middleware.gzip import GZipMiddleware
from fastapi.responses import JSONResponse, HTMLResponse, StreamingResponse
from sqlalchemy import text
from sqlalchemy.exc import OperationalError, DBAPIError
from sqlalchemy.orm import Session, joinedload

from app.config import settings
from app.database import get_db, engine
from app import models, schemas

# Initialize the FastAPI app
app = FastAPI(
    title="Backend Dulce Espera",
    description="API para la gestión de pedidos e insumos del sistema Dulce Espera.",
    version="1.0.0"
)

# Enable Gzip compression for fast network transfers
app.add_middleware(GZipMiddleware, minimum_size=500)

# Configure CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_origins_list,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# --- GLOBAL DATABASE ERROR HANDLERS ---
# If the MySQL database is unreachable or down, return 503 Service Unavailable

@app.exception_handler(OperationalError)
def db_operational_error_handler(request, exc: OperationalError):
    import logging
    logger = logging.getLogger("uvicorn")
    logger.exception("OperationalError in DB:")
    return JSONResponse(
        status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
        content={
            "detail": "No se pudo conectar a la base de datos MySQL. "
                      "Por favor, verifique la conexión y vuelva a intentarlo."
        }
    )

@app.exception_handler(DBAPIError)
def db_api_error_handler(request, exc: DBAPIError):
    return JSONResponse(
        status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
        content={
            "detail": "Error en la base de datos. Por favor, intente de nuevo más tarde."
        }
    )

# --- UTILIDADES ---

def esc(valor) -> str:
    """
    Escapa texto antes de interpolarlo en las plantillas HTML de los reportes.

    Los reportes se construyen con f-strings, así que cualquier texto que no
    venga de este archivo puede cerrar una etiqueta e inyectar marcado. El
    campo 'solicitante' lo escribe quien crea el pedido y los nombres del
    catálogo llegan desde FileMaker, de modo que ninguno de los dos es de
    confianza. Devuelve cadena vacía para None para no imprimir "None".
    """
    return html.escape(str(valor)) if valor is not None else ""


# --- ENDPOINTS ---

@app.get("/health", status_code=status.HTTP_200_OK, tags=["Control de Salud"])
def health_check(db: Session = Depends(get_db)):
    """
    Endpoint de salud rápido para verificar la API y la conexión a la base de datos.
    """
    try:
        # Simple query to check DB availability
        db.execute(text("SELECT 1"))
        return {"status": "ok", "database": "connected"}
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="El servicio está activo pero la base de datos no responde."
        )


@app.get("/insumos", response_model=List[schemas.InsumoOut], tags=["Insumos"])
def obtener_insumos(db: Session = Depends(get_db)):
    """
    Devuelve la lista de insumos activos cargados en el sistema.
    Retorna una lista vacía si aún no se han exportado insumos desde FileMaker.
    """
    insumos = db.query(models.Insumo).filter(
        models.Insumo.activo == 1
    ).order_by(models.Insumo.nombre.asc()).all()
    return insumos


@app.post("/pedidos", response_model=schemas.PedidoOut, status_code=status.HTTP_201_CREATED, tags=["Pedidos"])
def crear_pedido(pedido_in: schemas.PedidoCreate, db: Session = Depends(get_db)):
    """
    Crea un nuevo pedido de insumos.
    - Genera un UUID para el pedido.
    - Valida que todos los insumos solicitados existan y estén activos en la base de datos.
    - Valida que las cantidades sean mayores que cero (gestionado por el esquema de Pydantic).
    - Inserta los detalles de pedido correspondientes en una sola transacción.
    """
    pedido_uuid = str(uuid.uuid4())
    ahora = datetime.now()
    import logging
    logger = logging.getLogger("uvicorn")
    logger.info(f"RECIBIDO CREAR PEDIDO: {pedido_in.model_dump()}")

    # 1. Verificar insumos
    insumo_ids = {linea.insumo_id_publico for linea in pedido_in.lineas}
    
    # Query active insumos matching the requested IDs
    insumos_activos = db.query(models.Insumo).filter(
        models.Insumo.id_publico.in_(insumo_ids),
        models.Insumo.activo == 1
    ).all()
    
    insumos_activos_dict = {i.id_publico: i for i in insumos_activos}

    # Validar que todos los insumos del pedido realmente existan y estén activos
    for linea in pedido_in.lineas:
        if linea.insumo_id_publico not in insumos_activos_dict:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"El insumo con ID público '{linea.insumo_id_publico}' no existe o no está activo en el catálogo."
            )

    # 2. Iniciar transacción e insertar datos
    try:
        # Crear pedido
        db_pedido = models.Pedido(
            id_publico=pedido_uuid,
            solicitante=pedido_in.solicitante,
            fecha_solicitud=ahora,
            estado="pendiente",
            fecha_estado=ahora,
            motivo=pedido_in.motivo
        )
        db.add(db_pedido)

        # Crear líneas de detalle
        for linea in pedido_in.lineas:
            detalle_uuid = str(uuid.uuid4())
            db_detalle = models.DetallePedido(
                id_publico=detalle_uuid,
                pedido_id_publico=pedido_uuid,
                insumo_id_publico=linea.insumo_id_publico,
                cantidad=linea.cantidad
            )
            db.add(db_detalle)

        db.commit()
    except Exception as e:
        db.rollback()
        raise e

    # 3. Retornar el pedido con todas sus relaciones cargadas en una sola consulta
    pedido_creado = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.id_publico == pedido_uuid).first()

    if not pedido_creado:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Error al recuperar el pedido recién creado."
        )

    return pedido_creado


@app.get("/pedidos", response_model=List[schemas.PedidoOut], tags=["Pedidos"])
def consultar_pedidos(solicitante: str, db: Session = Depends(get_db)):
    """
    Devuelve la lista de pedidos realizados por un solicitante específico,
    incluyendo sus correspondientes líneas de detalle e información del insumo.
    """
    if not solicitante.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El parámetro 'solicitante' no puede estar vacío."
        )

    # Obtenemos los pedidos del solicitante, ordenados por fecha de solicitud descendente.
    # Usamos joinedload para evitar el problema de N+1 consultas.
    pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(
        models.Pedido.solicitante == solicitante
    ).order_by(
        models.Pedido.fecha_solicitud.desc()
    ).all()

    return pedidos


@app.get("/coordinadores", response_model=List[schemas.CoordinadorOut], tags=["Coordinadores"])
def obtener_coordinadores(db: Session = Depends(get_db)):
    """
    Devuelve la lista de coordinadores activos registrados en el sistema,
    para que la cocina pueda seleccionar a quién enviarle el pedido.
    """
    return db.query(models.Coordinador).filter(models.Coordinador.activo == 1).all()


# --- AUTENTICACIÓN ---

@app.post("/login", tags=["Autenticación"])
def login(login_in: schemas.LoginRequest, db: Session = Depends(get_db)):
    """
    Autenticación simple contra la tabla 'usuarios' en MySQL.
    Retorna un token, nombre y rol del usuario autenticado.
    Los usuarios se gestionan directamente en phpMyAdmin.
    """
    usuario = db.query(models.Usuario).filter(
        models.Usuario.username == login_in.username,
        models.Usuario.activo == 1
    ).first()

    if not usuario or usuario.password != login_in.password:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Credenciales inválidas. Verifique usuario y contraseña."
        )

    token = hashlib.sha256(f"{usuario.username}{time.time()}".encode()).hexdigest()[:32]

    return {
        "status": "success",
        "token": token,
        "nombre": usuario.nombre_display,
        "username": usuario.username,
        "rol": usuario.rol
    }


# --- ENDPOINTS PARA FILEMAKER ---

@app.get("/api/pedidos/pendientes", tags=["FileMaker"])
def obtener_pedidos_pendientes_fm(db: Session = Depends(get_db)):
    """
    Devuelve un JSON estructurado de forma limpia con los pedidos que tengan estado = 'pendiente'.
    Cada pedido incluye su lista de detalles anidada, incluyendo detalles descriptivos de los insumos.
    """
    pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(
        models.Pedido.estado == "pendiente"
    ).order_by(
        models.Pedido.fecha_solicitud.asc()
    ).all()

    resultado = []
    for p in pedidos:
        detalles_lista = []
        for l in p.lineas:
            detalles_lista.append({
                "id_publico": l.id_publico,
                "insumo_id_publico": l.insumo_id_publico,
                "nombre_insumo": l.insumo.nombre if l.insumo else None,
                "categoria_insumo": l.insumo.categoria if l.insumo else None,
                "grupo_insumo": l.insumo.grupo if l.insumo else None,
                "presentacion_insumo": l.insumo.presentacion if l.insumo else None,
                "cantidad": float(l.cantidad) if l.cantidad else 0.0
            })
        resultado.append({
            "id_publico": p.id_publico,
            "solicitante": p.solicitante,
            "fecha_solicitud": p.fecha_solicitud.strftime("%Y-%m-%d %H:%M:%S") if p.fecha_solicitud else None,
            "motivo": p.motivo,
            "detalles": detalles_lista
        })

    return resultado


@app.get("/api/pedidos/detalles-pendientes-flat", tags=["FileMaker"])
def obtener_detalles_pendientes_flat_fm(db: Session = Depends(get_db)):
    """
    Devuelve una lista completamente plana (flat) con todas las líneas de detalles
    de los pedidos que están en estado = 'pendiente'.
    Esto simplifica el consumo desde FileMaker al evitar la necesidad de recorrer arrays aninados.
    """
    pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(
        models.Pedido.estado == "pendiente"
    ).order_by(
        models.Pedido.fecha_solicitud.asc()
    ).all()

    resultado = []
    for p in pedidos:
        for l in p.lineas:
            resultado.append({
                "pedido_id_publico": p.id_publico,
                "solicitante": p.solicitante,
                "fecha_solicitud": p.fecha_solicitud.strftime("%Y-%m-%d %H:%M:%S") if p.fecha_solicitud else None,
                "estado": p.estado,
                "motivo": p.motivo,
                "detalle_id_publico": l.id_publico,
                "insumo_id_publico": l.insumo_id_publico,
                "nombre_insumo": l.insumo.nombre if l.insumo else None,
                "categoria_insumo": l.insumo.categoria if l.insumo else None,
                "grupo_insumo": l.insumo.grupo if l.insumo else None,
                "presentacion_insumo": l.insumo.presentacion if l.insumo else None,
                "cantidad": float(l.cantidad) if l.cantidad else 0.0
            })

    return resultado



@app.patch("/api/pedidos/actualizar-estado", tags=["FileMaker"])
def actualizar_estado_pedido_fm(body: schemas.ActualizarEstadoRequest, db: Session = Depends(get_db)):
    """
    Actualiza el estado de un pedido en MySQL desde FileMaker.
    Recibe un JSON con id_publico y el nuevo estado.
    """
    pedido = db.query(models.Pedido).filter(
        models.Pedido.id_publico == body.id_publico
    ).first()

    if not pedido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Pedido con ID '{body.id_publico}' no encontrado."
        )

    estados_validos = ["pendiente", "aceptado", "rechazado", "en revision", "comprado", "entregado", "cancelado"]
    estado_normalizado = body.estado.lower().strip()

    if estado_normalizado not in estados_validos:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Estado '{body.estado}' no válido. Opciones: {', '.join(estados_validos)}"
        )

    estado_anterior = pedido.estado
    pedido.estado = estado_normalizado
    pedido.fecha_estado = datetime.now()

    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise e

    return {
        "status": "success",
        "id_publico": pedido.id_publico,
        "estado_anterior": estado_anterior,
        "estado_nuevo": pedido.estado,
        "fecha_estado": pedido.fecha_estado.strftime("%Y-%m-%d %H:%M:%S")
    }


@app.get("/api/pedidos-pendientes", tags=["FileMaker"])
def pedidos_pendientes(estado: str = "pendiente", db: Session = Depends(get_db)):
    """
    Devuelve los pedidos filtrados por estado en formato JSON plano.
    Diseñado para que FileMaker lo consuma con 'Insertar desde URL' (GET).

    Uso desde FileMaker:
      Insertar desde URL [ $respuesta ; "https://tu-api/api/pedidos-pendientes" ]
      Insertar desde URL [ $respuesta ; "https://tu-api/api/pedidos-pendientes?estado=aceptado" ]
    """
    pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(
        models.Pedido.estado == estado.lower().strip()
    ).order_by(
        models.Pedido.fecha_solicitud.asc()
    ).all()

    resultado = []
    for p in pedidos:
        lineas_plano = []
        for l in p.lineas:
            lineas_plano.append({
                "insumo_id_publico": l.insumo_id_publico,
                "nombre_insumo": l.insumo.nombre if l.insumo else None,
                "categoria": l.insumo.categoria if l.insumo else None,
                "presentacion": l.insumo.presentacion if l.insumo else None,
                "cantidad": float(l.cantidad) if l.cantidad else 0
            })
        resultado.append({
            "id_publico": p.id_publico,
            "solicitante": p.solicitante,
            "fecha_solicitud": p.fecha_solicitud.isoformat() if p.fecha_solicitud else None,
            "estado": p.estado,
            "fecha_estado": p.fecha_estado.isoformat() if p.fecha_estado else None,
            "motivo": p.motivo,
            "total_lineas": len(lineas_plano),
            "lineas": lineas_plano
        })

    return {"pedidos": resultado, "total": len(resultado)}


@app.patch("/pedidos/actualizar-estado", tags=["FileMaker"])
def actualizar_estado_pedido(body: schemas.ActualizarEstadoRequest, db: Session = Depends(get_db)):
    """
    Actualiza el estado de un pedido existente.
    Diseñado para que FileMaker envíe un PATCH cuando el administrador
    aprueba o rechaza un pedido.

    Uso desde FileMaker:
      Establecer variable [ $json ; JSONSetElement("{}"; ["id_publico"; $id; JSONString]; ["estado"; "aceptado"; JSONString]) ]
      Establecer variable [ $curl ; "-X PATCH -H \"Content-Type: application/json\" -d " & Quote($json) ]
      Insertar desde URL [ $respuesta ; "https://tu-api/pedidos/actualizar-estado" ; $curl ]

    Estados válidos: pendiente, aceptado, rechazado, en revision, comprado, entregado, cancelado
    """
    pedido = db.query(models.Pedido).filter(
        models.Pedido.id_publico == body.id_publico
    ).first()

    if not pedido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Pedido con ID '{body.id_publico}' no encontrado."
        )

    estados_validos = ["pendiente", "aceptado", "rechazado", "en revision", "comprado", "entregado", "cancelado"]
    estado_normalizado = body.estado.lower().strip()

    if estado_normalizado not in estados_validos:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Estado '{body.estado}' no válido. Opciones: {', '.join(estados_validos)}"
        )

    estado_anterior = pedido.estado
    pedido.estado = estado_normalizado
    pedido.fecha_estado = datetime.now()

    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise e

    return {
        "status": "success",
        "id_publico": pedido.id_publico,
        "estado_anterior": estado_anterior,
        "estado_nuevo": pedido.estado,
        "fecha_estado": pedido.fecha_estado.isoformat()
    }


@app.get("/pedidos/todos", response_model=List[schemas.PedidoOut], tags=["Pedidos"])
def obtener_todos_pedidos(db: Session = Depends(get_db)):
    """
    Devuelve TODOS los pedidos del sistema sin filtrar por solicitante.
    Incluye líneas de detalle e información del insumo.
    La PWA usa este endpoint para mostrar el estado global de pedidos.
    """
    pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).order_by(
        models.Pedido.fecha_solicitud.desc()
    ).all()

    return pedidos


@app.get("/pedidos/{id_publico}/reporte", response_class=HTMLResponse, tags=["Pedidos"])
def ver_reporte_pedido(id_publico: str, db: Session = Depends(get_db)):
    """
    Genera una página HTML imprimible y elegante con el reporte completo del pedido,
    diseñada para ser enviada por WhatsApp y abierta en cualquier dispositivo.
    """
    pedido = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.id_publico == id_publico).first()

    if not pedido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pedido no encontrado"
        )

    fecha_str = pedido.fecha_solicitud.strftime("%Y-%m-%d %H:%M") if pedido.fecha_solicitud else "N/A"
    
    # Generar filas de la tabla
    lineas_html = ""
    for idx, linea in enumerate(pedido.lineas, 1):
        nombre = esc(linea.insumo.nombre) if linea.insumo else "Insumo sin nombre"
        presentacion = esc(linea.insumo.presentacion) if linea.insumo else "Unidades"
        categoria = esc(linea.insumo.categoria) if linea.insumo else "Otros"
        cantidad_val = float(linea.cantidad) if linea.cantidad else 0.0
        
        lineas_html += f"""
        <tr style="border-bottom: 1px solid #e2e8f0;">
            <td style="padding: 12px 10px; text-align: center; font-weight: bold; color: #64748b;">{idx}</td>
            <td style="padding: 12px 10px; font-weight: bold; color: #0f172a; text-align: left;">{nombre}</td>
            <td style="padding: 12px 10px; color: #475569; font-size: 13px; text-align: left;">{categoria}</td>
            <td style="padding: 12px 10px; text-align: right; font-weight: 800; color: #006156; font-size: 15px;">{cantidad_val:.2f}</td>
            <td style="padding: 12px 10px; color: #475569; font-weight: 500; font-size: 13px; text-align: left;">{presentacion}</td>
        </tr>
        """

    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Reporte de Pedido - Dulce Espera</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: #0f172a;
            background-color: #f8fafc;
            margin: 0;
            padding: 20px;
        }}
        .container {{
            max-width: 800px;
            margin: 0 auto;
            background-color: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 16px;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1), 0 2px 4px -1px rgba(0, 0, 0, 0.06);
            padding: 40px;
            box-sizing: border-box;
        }}
        .header {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            border-bottom: 3px solid #006156;
            padding-bottom: 24px;
            margin-bottom: 30px;
        }}
        .brand {{
            display: flex;
            align-items: center;
            gap: 16px;
        }}
        .brand-title {{
            font-size: 24px;
            font-weight: 800;
            color: #006156;
            margin: 0;
            letter-spacing: -0.025em;
        }}
        .brand-subtitle {{
            font-size: 13px;
            color: #39ADA3;
            font-weight: 700;
            margin-top: 4px;
        }}
        .meta {{
            text-align: right;
            font-size: 13px;
            color: #475569;
            line-height: 1.5;
        }}
        .meta strong {{
            color: #0f172a;
        }}
        .details-table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 20px;
            margin-bottom: 40px;
        }}
        .details-table th {{
            background-color: #f8fafc;
            color: #006156;
            font-weight: 700;
            text-transform: uppercase;
            font-size: 11px;
            letter-spacing: 0.05em;
            padding: 12px 10px;
            border-bottom: 2px solid #cbd5e1;
        }}
        .print-btn {{
            display: inline-flex;
            align-items: center;
            justify-content: center;
            background-color: #006156;
            color: #ffffff;
            border: none;
            border-radius: 8px;
            padding: 12px 24px;
            font-size: 14px;
            font-weight: 700;
            cursor: pointer;
            transition: background-color 0.15s ease;
            margin-bottom: 20px;
            gap: 8px;
            text-decoration: none;
        }}
        .print-btn:hover {{
            background-color: #004d45;
        }}
        .no-print {{
            display: flex;
            justify-content: flex-end;
            max-width: 800px;
            margin: 0 auto;
        }}
        .footer {{
            border-top: 1px solid #e2e8f0;
            padding-top: 20px;
            text-align: center;
            font-size: 12px;
            color: #64748b;
            margin-top: 50px;
            line-height: 1.5;
        }}
        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .container {{
                border: none;
                box-shadow: none;
                padding: 0;
                max-width: 100%;
            }}
            .no-print {{
                display: none !important;
            }}
        }}
    </style>
</head>
<body>
    <div class="no-print">
        <button class="print-btn" onclick="window.print()">
            <svg style="width: 18px; height: 18px;" fill="none" stroke="currentColor" viewBox="0 0 24 24" xmlns="http://www.w3.org/2000/svg">
                <path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M17 17h2a2 2 0 002-2v-4a2 2 0 00-2-2H5a2 2 0 00-2 2v4a2 2 0 002 2h2m2 4h10a2 2 0 002-2v-4a2 2 0 00-2-2H9a2 2 0 00-2 2v4a2 2 0 002 2zm8-12V5a2 2 0 00-2-2H9a2 2 0 00-2 2v4h10z"></path>
            </svg>
            Imprimir / Guardar PDF
        </button>
    </div>
    <div class="container">
        <div class="header">
            <div class="brand">
                <img src="https://dulce-espera-inventario.vercel.app/logo.svg" alt="Logo" style="width: 50px; height: 50px; object-fit: contain;" onerror="this.style.display='none'">
                <div>
                    <h1 class="brand-title" style="margin: 0; line-height: 1;">DULCE ESPERA</h1>
                </div>
            </div>
            <div class="meta">
                <div><strong>ID Solicitud:</strong> {pedido.id_publico[:8].upper()}</div>
                <div style="font-size: 10px; color: #94a3b8; margin-bottom: 4px;">UUID: {pedido.id_publico}</div>
                <div><strong>Fecha:</strong> {fecha_str}</div>
                <div><strong>Solicitante:</strong> {esc(pedido.solicitante)}</div>
                <div><strong>Estado:</strong> <span style="text-transform: uppercase; font-weight: 800; color: #b45309;">{esc(pedido.estado)}</span></div>
            </div>
        </div>

        <h2 style="font-size: 16px; margin-bottom: 20px; font-weight: 700; color: #1e293b; border-left: 4px solid #006156; padding-left: 10px; text-align: left;">
            Detalle del Pedido de Insumos
        </h2>

        <table class="details-table">
            <thead>
                <tr>
                    <th style="width: 50px; text-align: center;">Item</th>
                    <th style="text-align: left;">Descripción Insumo</th>
                    <th style="text-align: left;">Categoría</th>
                    <th style="text-align: right; width: 100px;">Cantidad</th>
                    <th style="text-align: left; width: 120px;">Presentación</th>
                </tr>
            </thead>
            <tbody>
                {lineas_html}
            </tbody>
        </table>

        <div style="margin-top: 80px; display: flex; justify-content: space-around;">
            <div style="text-align: center; width: 220px; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 12px; color: #475569; font-weight: 600;">
                Firma Responsable
            </div>
            <div style="text-align: center; width: 220px; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 12px; color: #475569; font-weight: 600;">
                Firma Autorización
            </div>
        </div>

        <div class="footer">
            &copy; {datetime.now().year} Dulce Espera.
        </div>
    </div>
</body>
</html>"""
    return HTMLResponse(content=html_content)


def resolver_canal_compra(linea) -> str:
    """Obtiene el canal de compra del insumo: Mercado, Super Mercado, Proveedor u Otros."""
    if linea.insumo and linea.insumo.grupo:
        g = linea.insumo.grupo.strip()
        g_lower = g.lower()
        if g_lower == "mercado":
            return "Mercado"
        elif g_lower in ("super mercado", "super"):
            return "Super Mercado"
        elif g_lower == "proveedor":
            return "Proveedor"
        return g
    cat = (linea.insumo.categoria or "").lower() if linea.insumo else ""
    if any(kw in cat for kw in ["verdura", "fruta", "carne", "proteina", "pollo", "pescado", "fresco", "huevo"]):
        return "Mercado"
    return "Super Mercado"


@app.get("/pedidos/{id_publico}/reporte/pdf", tags=["Pedidos"])
def ver_reporte_pedido_pdf(id_publico: str, db: Session = Depends(get_db)):
    """
    Genera y descarga un archivo PDF del reporte del pedido de forma directa en base al HTML,
    estructurado y clasificado por canales de compra.
    """
    from xhtml2pdf import pisa
    import io
    from fastapi.responses import StreamingResponse

    pedido = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.id_publico == id_publico).first()

    if not pedido:
        raise HTTPException(
            status_code=404,
            detail="Pedido no encontrado"
        )

    fecha_str = pedido.fecha_solicitud.strftime("%Y-%m-%d %H:%M") if pedido.fecha_solicitud else "N/A"
    
    lineas_mercado = []
    lineas_super = []
    lineas_prov = []
    lineas_otros = []

    for linea in pedido.lineas:
        item_data = {
            "nombre": esc(linea.insumo.nombre) if linea.insumo else "Insumo sin nombre",
            "presentacion": esc(linea.insumo.presentacion) if linea.insumo else "Unidades",
            "categoria": esc(linea.insumo.categoria) if linea.insumo else "Otros",
            "cantidad": float(linea.cantidad) if linea.cantidad else 0.0,
            "id_insumo": linea.insumo_id_publico or "N/A"
        }
        canal = resolver_canal_compra(linea)
        if canal == "Mercado":
            lineas_mercado.append(item_data)
        elif canal == "Super Mercado":
            lineas_super.append(item_data)
        elif canal == "Proveedor":
            lineas_prov.append(item_data)
        else:
            lineas_otros.append(item_data)

    def render_tabla_canal_pdf(titulo, subtitulo, lineas):
        if not lineas:
            return ""
        filas = ""
        total_canal = sum(x["cantidad"] for x in lineas)
        for idx, it in enumerate(lineas, 1):
            filas += f"""
            <tr>
                <td style="padding: 5px 4px; text-align: center; border-bottom: 1px solid #e2e8f0; font-size: 9px; color: #64748b;">{idx}</td>
                <td style="padding: 5px 8px; font-weight: bold; text-align: left; border-bottom: 1px solid #e2e8f0; font-size: 10px; color: #0f172a;">{it['nombre']}</td>
                <td style="padding: 5px 8px; text-align: left; border-bottom: 1px solid #e2e8f0; font-size: 9px; color: #475569;">{it['categoria']}</td>
                <td style="padding: 5px 8px; text-align: right; font-weight: bold; color: #006156; border-bottom: 1px solid #e2e8f0; font-size: 10px;">{it['cantidad']:.2f}</td>
                <td style="padding: 5px 8px; text-align: left; border-bottom: 1px solid #e2e8f0; font-size: 9px; color: #475569;">{it['presentacion']}</td>
                <td style="padding: 5px 8px; text-align: center; border-bottom: 1px solid #e2e8f0; font-size: 9px; color: #cbd5e1;">[ &nbsp; ]</td>
            </tr>
            """
        return f"""
        <div style="margin-top: 14px; margin-bottom: 4px; font-size: 10px; font-weight: bold; color: #006156; border-left: 3px solid #006156; padding-left: 6px;">
            {titulo.upper()} &mdash; <span style="font-size: 9px; font-weight: normal; color: #64748b;">{subtitulo}</span>
        </div>
        <table class="details-table" style="width: 100%; border-collapse: collapse; margin-bottom: 10px;">
            <thead>
                <tr style="background-color: #f1f5f9;">
                    <th style="width: 25px; text-align: center; font-size: 9px; padding: 5px; color: #1e293b;">N°</th>
                    <th style="text-align: left; font-size: 9px; padding: 5px 8px; color: #1e293b;">Descripción Insumo</th>
                    <th style="text-align: left; width: 130px; font-size: 9px; padding: 5px 8px; color: #1e293b;">Categoría</th>
                    <th style="text-align: right; width: 70px; font-size: 9px; padding: 5px 8px; color: #1e293b;">Cant.</th>
                    <th style="text-align: left; width: 80px; font-size: 9px; padding: 5px 8px; color: #1e293b;">Unidad</th>
                    <th style="text-align: center; width: 50px; font-size: 9px; padding: 5px; color: #1e293b;">Check</th>
                </tr>
            </thead>
            <tbody>
                {filas}
                <tr style="background-color: #f8fafc;">
                    <td colspan="3" style="text-align: right; font-weight: bold; font-size: 9px; padding: 5px 8px; color: #0f172a;">SUBTOTAL {titulo.upper()}:</td>
                    <td style="text-align: right; font-weight: bold; font-size: 10px; padding: 5px 8px; color: #006156;">{total_canal:.2f}</td>
                    <td colspan="2"></td>
                </tr>
            </tbody>
        </table>
        """

    tablas_html = ""
    tablas_html += render_tabla_canal_pdf("Plaza de Mercado", "Perecederos, Carnes, Frutas y Verduras frescas", lineas_mercado)
    tablas_html += render_tabla_canal_pdf("Supermercado y Abarrotes", "Secos, Lácteos industriales, Granos y Limpieza", lineas_super)
    tablas_html += render_tabla_canal_pdf("Proveedores Directos", "Distribuidoras, Panadería, Kéfir y Especiales", lineas_prov)
    if lineas_otros:
        tablas_html += render_tabla_canal_pdf("Otros Insumos y Servicios", "Descartables y consumos varios", lineas_otros)

    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>Reporte de Pedido - Dulce Espera</title>
    <style>
        @page {{
            size: letter;
            margin: 20mm 20mm 20mm 20mm;
        }}
        body {{
            font-family: Helvetica, Arial, sans-serif;
            color: #0f172a;
            font-size: 11px;
            line-height: 1.4;
        }}
        .header-table {{
            width: 100%;
            margin-bottom: 15px;
            border-bottom: 2px solid #006156;
            padding-bottom: 10px;
        }}
        .brand-title {{
            font-size: 18px;
            font-weight: bold;
            color: #006156;
        }}
        .brand-subtitle {{
            font-size: 10px;
            color: #39ADA3;
            font-weight: bold;
        }}
        .meta-text {{
            font-size: 10px;
            color: #475569;
            text-align: right;
            line-height: 1.4;
        }}
        .info-table {{
            width: 100%;
            margin-bottom: 15px;
        }}
        .section-title {{
            font-size: 11px;
            font-weight: bold;
            color: #006156;
            border-left: 3px solid #39ADA3;
            padding-left: 8px;
            margin-bottom: 10px;
        }}
        .details-table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 25px;
        }}
        .details-table th {{
            background-color: #f8fafc;
            border-bottom: 2px solid #006156;
            color: #006156;
            font-weight: bold;
            font-size: 10px;
            padding: 8px 10px;
        }}
        .details-table td {{
            padding: 8px 10px;
            font-size: 10px;
        }}
        .footer {{
            border-top: 1px solid #cbd5e1;
            padding-top: 15px;
            text-align: center;
            font-size: 9px;
            color: #64748b;
            margin-top: 30px;
        }}
        .signature-table {{
            width: 100%;
            margin-top: 40px;
        }}
        .signature-line {{
            border-top: 1px solid #cbd5e1;
            text-align: center;
            font-size: 10px;
            color: #475569;
            font-weight: bold;
            padding-top: 5px;
        }}
    </style>
</head>
<body>
    <table class="header-table">
        <tr>
            <td style="width: 55px; vertical-align: middle;">
                <img src="https://dulce-espera-inventario.vercel.app/icon-192.png" style="width: 45px; height: 45px;" />
            </td>
            <td style="vertical-align: middle;">
                <div class="brand-title" style="line-height: 1;">DULCE ESPERA</div>
            </td>
            <td class="meta-text" style="vertical-align: middle;">
                <strong>N° LISTA:</strong> {pedido.id_publico[:8].upper()}<br>
                <strong>FECHA:</strong> {fecha_str}<br>
                <strong>ESTADO:</strong> <span style="color: #006156; font-weight: bold;">{pedido.estado.upper()}</span>
            </td>
        </tr>
    </table>

    <table class="info-table">
        <tr>
            <td style="font-size: 10px; color: #475569; line-height: 1.4;">
                <strong>Solicitado por:</strong> {esc(pedido.solicitante)}
            </td>
            <td style="font-size: 10px; color: #475569; text-align: right; line-height: 1.4; vertical-align: top;">
                &nbsp;
            </td>
        </tr>
    </table>

    {tablas_html}

    <table class="signature-table">
        <tr>
            <td style="width: 45%;">
                <div class="signature-line">Firma Solicitante Cocina</div>
            </td>
            <td style="width: 10%;">&nbsp;</td>
            <td style="width: 45%;">
                <div class="signature-line">Firma Autorización</div>
            </td>
        </tr>
    </table>

    <div class="footer">
        &copy; {datetime.now().year} Dulce Espera.
    </div>
</body>
</html>"""

    pdf_buffer = io.BytesIO()
    pisa_status = pisa.CreatePDF(html_content, dest=pdf_buffer)
    if pisa_status.err:
        raise HTTPException(
            status_code=500,
            detail="Error al generar el PDF"
        )
    pdf_buffer.seek(0)
    filename = f"Pedido_{pedido.id_publico[:8].upper()}.pdf"
    return StreamingResponse(
        pdf_buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )




def render_tabla_canal(titulo: str, subtitulo: str, color_hex: str, lineas: list) -> str:
    if not lineas:
        return f"""
        <div style="padding: 12px 16px; border-left: 3px solid #cbd5e1; background-color: #f8fafc; color: #64748b; font-size: 12px; font-weight: 500; margin-bottom: 24px; font-style: italic; border-radius: 4px;">
            No se registraron requerimientos para el canal de {titulo}.
        </div>
        """

    lineas_ordenadas = sorted(lineas, key=lambda x: (x["categoria"] or "", x["nombre"] or ""))
    total_cant = sum(x["cantidad"] for x in lineas_ordenadas)

    filas = ""
    for idx, item in enumerate(lineas_ordenadas, 1):
        filas += f"""
        <tr style="border-bottom: 1px solid #e2e8f0; page-break-inside: avoid;">
            <td style="padding: 9px 8px; text-align: center; font-weight: bold; color: #64748b; font-size: 12px;">{idx}</td>
            <td style="padding: 9px 8px; color: #64748b; font-size: 11px; font-family: monospace;">{item['id_insumo']}</td>
            <td style="padding: 9px 8px; font-weight: 700; color: #0f172a; text-align: left; font-size: 13px;">{item['nombre']}</td>
            <td style="padding: 9px 8px; color: #475569; font-size: 12px; text-align: left;"><span style="background: #f1f5f9; padding: 2px 7px; border-radius: 4px; font-weight: 600; font-size: 11px;">{item['categoria']}</span></td>
            <td style="padding: 9px 8px; text-align: right; font-weight: 800; color: {color_hex}; font-size: 14px;">{item['cantidad']:.2f}</td>
            <td style="padding: 9px 8px; color: #475569; font-weight: 600; font-size: 12px; text-align: left;">{item['presentacion']}</td>
            <td style="padding: 9px 8px; text-align: center; vertical-align: middle;">
                <div style="width: 16px; height: 16px; border: 2px solid #cbd5e1; border-radius: 4px; margin: 0 auto;"></div>
            </td>
        </tr>
        """

    return f"""
    <div style="margin-bottom: 28px; page-break-inside: avoid;">
        <div style="display: flex; justify-content: space-between; align-items: baseline; margin-bottom: 10px; border-left: 4px solid {color_hex}; padding-left: 10px;">
            <div>
                <h3 style="font-size: 13px; font-weight: 900; color: #0f172a; margin: 0; text-transform: uppercase; letter-spacing: 0.04em;">
                    {titulo}
                </h3>
                <span style="font-size: 11px; color: #64748b; font-weight: 500;">{subtitulo}</span>
            </div>
            <div style="font-size: 11px; color: #64748b; font-weight: 700;">
                {len(lineas_ordenadas)} insumo{'s' if len(lineas_ordenadas) != 1 else ''} | {total_cant:.2f} unidades
            </div>
        </div>
        <table style="width: 100%; border-collapse: collapse; margin-bottom: 8px;">
            <thead>
                <tr style="background-color: #f8fafc; border-bottom: 2px solid #cbd5e1;">
                    <th style="padding: 8px 8px; width: 40px; text-align: center; font-size: 11px; color: #475569; text-transform: uppercase;">N°</th>
                    <th style="padding: 8px 8px; width: 75px; text-align: left; font-size: 11px; color: #475569; text-transform: uppercase;">Código</th>
                    <th style="padding: 8px 8px; text-align: left; font-size: 11px; color: #475569; text-transform: uppercase;">Descripción Insumo</th>
                    <th style="padding: 8px 8px; text-align: left; width: 140px; font-size: 11px; color: #475569; text-transform: uppercase;">Categoría</th>
                    <th style="padding: 8px 8px; text-align: right; width: 95px; font-size: 11px; color: #475569; text-transform: uppercase;">Cant. Sol.</th>
                    <th style="padding: 8px 8px; text-align: left; width: 100px; font-size: 11px; color: #475569; text-transform: uppercase;">Presentación</th>
                    <th style="padding: 8px 8px; width: 75px; text-align: center; font-size: 11px; color: #475569; text-transform: uppercase;">Comprado</th>
                </tr>
            </thead>
            <tbody>
                {filas}
            </tbody>
        </table>
    </div>
    """


@app.get("/pedidos/{id_publico}/reporte-admin", response_class=HTMLResponse, tags=["Pedidos"])
def ver_reporte_pedido_admin(id_publico: str, db: Session = Depends(get_db)):
    """
    Genera un informe administrativo profesional del pedido, estructurado y clasificado
    por canales de compra reales: Mercado (Plaza/Perecederos), Supermercado y Proveedores Directos.
    """
    pedido = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.id_publico == id_publico).first()

    if not pedido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pedido no encontrado"
        )

    fecha_str = pedido.fecha_solicitud.strftime("%Y-%m-%d %H:%M") if pedido.fecha_solicitud else "N/A"
    
    lineas_mercado = []
    lineas_super = []
    lineas_prov = []
    lineas_otros = []

    for linea in pedido.lineas:
        item_data = {
            "nombre": esc(linea.insumo.nombre) if linea.insumo else "Insumo sin nombre",
            "presentacion": esc(linea.insumo.presentacion) if linea.insumo else "Unidades",
            "categoria": esc(linea.insumo.categoria) if linea.insumo else "Otros",
            "cantidad": float(linea.cantidad) if linea.cantidad else 0.0,
            "id_insumo": linea.insumo_id_publico or "N/A"
        }
        canal = resolver_canal_compra(linea)
        if canal == "Mercado":
            lineas_mercado.append(item_data)
        elif canal == "Super Mercado":
            lineas_super.append(item_data)
        elif canal == "Proveedor":
            lineas_prov.append(item_data)
        else:
            lineas_otros.append(item_data)

    html_secciones = ""
    html_secciones += render_tabla_canal("Grupo A: Plaza de Mercado", "Perecederos, Carnes, Frutas y Verduras frescas", "#b45309", lineas_mercado)
    html_secciones += render_tabla_canal("Grupo B: Supermercado y Abarrotes", "Secos, Lácteos industriales, Granos y Limpieza", "#006156", lineas_super)
    html_secciones += render_tabla_canal("Grupo C: Proveedores Directos", "Distribuidoras, Panadería, Kéfir y Especiales", "#4338ca", lineas_prov)
    if lineas_otros:
        html_secciones += render_tabla_canal("Grupo D: Otros Insumos y Servicios", "Descartables, envases y consumos varios", "#475569", lineas_otros)

    chips_resumen = f"""
    <div style="display: flex; gap: 14px; margin-bottom: 25px; border-bottom: 2px solid #cbd5e1; padding-bottom: 14px; flex-wrap: wrap;">
        <div style="background: #f8fafc; border: 1px solid #e2e8f0; border-radius: 8px; padding: 7px 14px; font-size: 11px; font-weight: 700; color: #475569;">
            TOTAL PEDIDO: <span style="font-size: 14px; color: #0f172a; margin-left: 4px;">{len(pedido.lineas)}</span>
        </div>
        <div style="background: #fef3c7; border: 1px solid #fde68a; border-radius: 8px; padding: 7px 14px; font-size: 11px; font-weight: 700; color: #92400e;">
            MERCADO: <span style="font-size: 14px; color: #b45309; margin-left: 4px;">{len(lineas_mercado)}</span>
        </div>
        <div style="background: #e6f0ef; border: 1px solid #b2d8d4; border-radius: 8px; padding: 7px 14px; font-size: 11px; font-weight: 700; color: #006156;">
            SUPERMERCADO: <span style="font-size: 14px; color: #006156; margin-left: 4px;">{len(lineas_super)}</span>
        </div>
        <div style="background: #eef2ff; border: 1px solid #c7d2fe; border-radius: 8px; padding: 7px 14px; font-size: 11px; font-weight: 700; color: #3730a3;">
            PROVEEDORES: <span style="font-size: 14px; color: #4338ca; margin-left: 4px;">{len(lineas_prov)}</span>
        </div>
        {f'<div style="background: #f1f5f9; border: 1px solid #cbd5e1; border-radius: 8px; padding: 7px 14px; font-size: 11px; font-weight: 700; color: #475569;">OTROS: <span style="font-size: 14px; color: #334155; margin-left: 4px;">{len(lineas_otros)}</span></div>' if lineas_otros else ''}
    </div>
    """

    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Informe de Compras por Canal - Dulce Espera</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: #0f172a;
            background-color: #f1f5f9;
            margin: 0;
            padding: 30px 20px;
        }}
        .container {{
            max-width: 900px;
            margin: 0 auto;
            background-color: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 20px;
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.05), 0 4px 6px -2px rgba(0, 0, 0, 0.02);
            padding: 45px;
        }}
        .header {{
            display: flex;
            justify-content: space-between;
            align-items: flex-start;
            border-bottom: 2px solid #cbd5e1;
            padding-bottom: 24px;
            margin-bottom: 30px;
        }}
        .brand {{
            display: flex;
            align-items: center;
            gap: 16px;
        }}
        .brand-title {{
            font-size: 26px;
            font-weight: 900;
            color: #006156;
            margin: 0;
            letter-spacing: -0.03em;
        }}
        .brand-subtitle {{
            font-size: 12px;
            color: #475569;
            font-weight: 700;
            margin-top: 3px;
            text-transform: uppercase;
            letter-spacing: 0.05em;
        }}
        .doc-type {{
            display: inline-block;
            background-color: #e6f0ef;
            color: #006156;
            font-size: 11px;
            font-weight: 800;
            padding: 4px 10px;
            border-radius: 6px;
            margin-top: 8px;
            text-transform: uppercase;
            letter-spacing: 0.05em;
        }}
        .meta {{
            text-align: right;
            font-size: 12px;
            color: #475569;
            line-height: 1.6;
        }}
        .print-btn {{
            background-color: #006156;
            color: white;
            border: none;
            padding: 10px 18px;
            border-radius: 10px;
            font-weight: bold;
            font-size: 13px;
            cursor: pointer;
            display: flex;
            align-items: center;
            gap: 8px;
            box-shadow: 0 4px 6px -1px rgba(0, 97, 86, 0.2);
            margin-bottom: 20px;
            margin-left: auto;
            margin-right: auto;
            max-width: 900px;
            transition: all 0.2s;
        }}
        .print-btn:hover {{
            background-color: #004d44;
        }}
        .footer {{
            border-top: 1px solid #e2e8f0;
            padding-top: 20px;
            text-align: center;
            font-size: 12px;
            color: #64748b;
            margin-top: 50px;
            line-height: 1.5;
        }}
        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .container {{
                border: none;
                box-shadow: none;
                padding: 0;
                max-width: 100%;
            }}
            .no-print {{
                display: none !important;
            }}
        }}
    </style>
</head>
<body>
    <div class="no-print">
        <button class="print-btn" onclick="window.print()">
            <svg style="width: 18px; height: 18px;" fill="none" stroke="currentColor" viewBox="0 0 24 24" xmlns="http://www.w3.org/2000/svg">
                <path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M17 17h2a2 2 0 002-2v-4a2 2 0 00-2-2H5a2 2 0 00-2 2v4a2 2 0 002 2h2m2 4h10a2 2 0 002-2v-4a2 2 0 00-2-2H9a2 2 0 00-2 2v4a2 2 0 002 2zm8-12V5a2 2 0 00-2-2H9a2 2 0 00-2 2v4h10z"></path>
            </svg>
            Imprimir / Guardar PDF
        </button>
    </div>
    <div class="container">
        <div class="header">
            <div class="brand">
                <img src="https://dulce-espera-inventario.vercel.app/logo.svg" alt="Logo" style="width: 75px; height: 75px; object-fit: contain;" onerror="this.style.display='none'">
                <div>
                    <h1 class="brand-title">DULCE ESPERA</h1>
                    <span class="doc-type">CONTROL DE COMPRAS POR CANALES</span>
                </div>
            </div>
            <div class="meta">
                <div><strong>ID Pedido:</strong> {pedido.id_publico[:8].upper()}</div>
                <div style="font-size: 10px; color: #94a3b8; margin-bottom: 2px;">UUID: {pedido.id_publico}</div>
                <div><strong>Fecha Pedido:</strong> {fecha_str}</div>
                <div><strong>Solicitante:</strong> {esc(pedido.solicitante)}</div>
                <div><strong>Estado Actual:</strong> <span style="text-transform: uppercase; font-weight: 800; color: #006156;">{esc(pedido.estado)}</span></div>
            </div>
        </div>

        {chips_resumen}

        {html_secciones}

        <div style="margin-top: 60px; display: flex; justify-content: space-between; gap: 30px; page-break-inside: avoid;">
            <div style="text-align: center; flex: 1; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 11px; color: #475569; font-weight: 600;">
                Firma Solicitante Cocina
            </div>
            <div style="text-align: center; flex: 1; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 11px; color: #475569; font-weight: 600;">
                Autorizado Administración / Compras
            </div>
            <div style="text-align: center; flex: 1; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 11px; color: #475569; font-weight: 600;">
                Recibido en Cocina (Control Físico)
            </div>
        </div>

        <div class="footer">
            Este informe está clasificado por canales de compra para optimizar los procesos logísticos.<br>
            &copy; {datetime.now().year} Dulce Espera.
        </div>
    </div>
</body>
</html>"""
    return HTMLResponse(content=html_content)


def generar_excel_pedido(pedido: models.Pedido):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado Compras"
    ws.views.sheetView[0].showGridLines = True

    TEAL_PRIMARY = "006156"
    DARK_TEXT = "0F172A"
    MUTED_TEXT = "64748B"
    BORDER_GRAY = "CBD5E1"

    font_banner = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    font_sub = Font(name="Segoe UI", size=9, bold=True, color="E2E8F0")
    font_lbl = Font(name="Segoe UI", size=9, bold=True, color=MUTED_TEXT)
    font_val = Font(name="Segoe UI", size=10, bold=True, color=DARK_TEXT)

    thin_border = Side(border_style="thin", color=BORDER_GRAY)
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # 1. Encabezado Institucional
    ws.merge_cells("A1:H1")
    ws["A1"] = "CLÍNICA MONTALVO — DULCE ESPERA"
    ws["A1"].font = font_banner
    ws["A1"].fill = PatternFill(start_color=TEAL_PRIMARY, end_color=TEAL_PRIMARY, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:H2")
    ws["A2"] = "ORDEN OFICIAL DE COMPRAS Y CONTROL DE SUMINISTROS"
    ws["A2"].font = font_sub
    ws["A2"].fill = PatternFill(start_color=TEAL_PRIMARY, end_color=TEAL_PRIMARY, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # 2. Metadatos
    fecha_str = pedido.fecha_solicitud.strftime("%Y-%m-%d %H:%M") if pedido.fecha_solicitud else "N/A"
    id_ped = f"#{pedido.id_publico[:8].upper()}" if pedido.id_publico else f"#{pedido.id}"

    ws["A4"] = "PEDIDO N°:"; ws["A4"].font = font_lbl
    ws["B4"] = id_ped; ws["B4"].font = font_val
    ws["C4"] = "FECHA:"; ws["C4"].font = font_lbl
    ws["D4"] = fecha_str; ws["D4"].font = font_val
    ws["E4"] = "SOLICITADO POR:"; ws["E4"].font = font_lbl
    ws["F4"] = pedido.solicitante or "N/A"; ws["F4"].font = font_val
    ws["G4"] = "ESTADO:"; ws["G4"].font = font_lbl
    ws["H4"] = (pedido.estado or "Pendiente").upper(); ws["H4"].font = font_val

    row_cursor = 5
    if pedido.motivo:
        ws[f"A{row_cursor}"] = "JUSTIFICACIÓN:"; ws[f"A{row_cursor}"].font = font_lbl
        ws.merge_cells(f"B{row_cursor}:H{row_cursor}")
        ws[f"B{row_cursor}"] = f'"{pedido.motivo}"'
        ws[f"B{row_cursor}"].font = Font(name="Segoe UI", size=10, italic=True, color="334155")
        ws.row_dimensions[row_cursor].height = 18
        row_cursor += 1

    # Agrupación por canal
    lineas_mercado = []
    lineas_super = []
    lineas_prov = []
    lineas_otros = []

    for linea in pedido.lineas:
        item_data = {
            "nombre": linea.insumo.nombre if linea.insumo else "Insumo sin nombre",
            "presentacion": linea.insumo.presentacion if linea.insumo else "Unidades",
            "categoria": linea.insumo.categoria if linea.insumo else "Otros",
            "cantidad": float(linea.cantidad) if linea.cantidad else 0.0,
            "id_insumo": linea.insumo_id_publico or "N/A"
        }
        canal = resolver_canal_compra(linea)
        if canal == "Mercado":
            lineas_mercado.append(item_data)
        elif canal == "Super Mercado":
            lineas_super.append(item_data)
        elif canal == "Proveedor":
            lineas_prov.append(item_data)
        else:
            lineas_otros.append(item_data)

    # 3. KPI Resumen
    row_kpi_lbl = row_cursor + 1
    row_kpi_val = row_kpi_lbl + 1

    kpis = [
        ("TOTAL LÍNEAS", str(len(pedido.lineas)), "A", "B"),
        ("MERCADO", str(len(lineas_mercado)), "C", "D"),
        ("SUPERMERCADO", str(len(lineas_super)), "E", "F"),
        ("PROVEEDORES", str(len(lineas_prov)), "G", "H")
    ]
    fill_kpi = PatternFill(start_color="E6F0EF", end_color="E6F0EF", fill_type="solid")
    font_kpi_lbl = Font(name="Segoe UI", size=8, bold=True, color="475569")
    font_kpi_num = Font(name="Segoe UI", size=13, bold=True, color=TEAL_PRIMARY)

    for label, val, c1, c2 in kpis:
        ws.merge_cells(f"{c1}{row_kpi_lbl}:{c2}{row_kpi_lbl}")
        ws.merge_cells(f"{c1}{row_kpi_val}:{c2}{row_kpi_val}")
        c_lbl = ws[f"{c1}{row_kpi_lbl}"]
        c_lbl.value = label
        c_lbl.font = font_kpi_lbl
        c_lbl.alignment = Alignment(horizontal="center", vertical="center")
        c_lbl.fill = fill_kpi

        c_v = ws[f"{c1}{row_kpi_val}"]
        c_v.value = val
        c_v.font = font_kpi_num
        c_v.alignment = Alignment(horizontal="center", vertical="center")
        c_v.fill = fill_kpi

    ws.row_dimensions[row_kpi_lbl].height = 16
    ws.row_dimensions[row_kpi_val].height = 22

    col_widths = {"A": 6, "B": 38, "C": 20, "D": 15, "E": 14, "F": 14, "G": 16, "H": 28}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    current_row = row_kpi_val + 2

    def escribir_tabla_canal(hoja, start_r, titulo, subtitulo, color_hex, lineas):
        hoja.merge_cells(f"A{start_r}:H{start_r}")
        c_sec = hoja[f"A{start_r}"]
        c_sec.value = f"{titulo.upper()} — {subtitulo}"
        c_sec.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        c_sec.fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
        c_sec.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        hoja.row_dimensions[start_r].height = 24

        th_r = start_r + 1
        headers = [
            ("A", "N°", Alignment(horizontal="center")),
            ("B", "Descripción del Insumo", Alignment(horizontal="left")),
            ("C", "Categoría", Alignment(horizontal="left")),
            ("D", "Cant. Solicitada", Alignment(horizontal="right")),
            ("E", "Unidad", Alignment(horizontal="center")),
            ("F", "Verificación [✓]", Alignment(horizontal="center")),
            ("G", "Cant. Recibida Real", Alignment(horizontal="right")),
            ("H", "Observaciones / Novedades", Alignment(horizontal="left"))
        ]
        for col, title, align in headers:
            c = hoja[f"{col}{th_r}"]
            c.value = title
            c.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            c.alignment = align
            c.border = border_cell
        hoja.row_dimensions[th_r].height = 20

        r_cursor = th_r + 1
        if not lineas:
            hoja.merge_cells(f"A{r_cursor}:H{r_cursor}")
            empty_c = hoja[f"A{r_cursor}"]
            empty_c.value = f"No se registraron requerimientos para {titulo}."
            empty_c.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
            empty_c.alignment = Alignment(horizontal="center", vertical="center")
            for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
                hoja[f"{col}{r_cursor}"].border = border_cell
            return r_cursor + 2

        lineas_ordenadas = sorted(lineas, key=lambda x: (x["categoria"] or "", x["nombre"] or ""))
        for idx, item in enumerate(lineas_ordenadas, 1):
            is_zebra = (idx % 2 == 0)
            row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if is_zebra else None

            hoja[f"A{r_cursor}"].value = idx
            hoja[f"A{r_cursor}"].alignment = Alignment(horizontal="center")
            hoja[f"A{r_cursor}"].font = Font(name="Segoe UI", size=9, color="64748B", bold=True)

            hoja[f"B{r_cursor}"].value = item["nombre"]
            hoja[f"B{r_cursor}"].alignment = Alignment(horizontal="left")
            hoja[f"B{r_cursor}"].font = Font(name="Segoe UI", size=10, color="0F172A", bold=True)

            hoja[f"C{r_cursor}"].value = item["categoria"]
            hoja[f"C{r_cursor}"].alignment = Alignment(horizontal="left")
            hoja[f"C{r_cursor}"].font = Font(name="Segoe UI", size=9, color="475569")

            hoja[f"D{r_cursor}"].value = item["cantidad"]
            hoja[f"D{r_cursor}"].number_format = "#,##0.00"
            hoja[f"D{r_cursor}"].alignment = Alignment(horizontal="right")
            hoja[f"D{r_cursor}"].font = Font(name="Segoe UI", size=10, bold=True, color="006156")

            hoja[f"E{r_cursor}"].value = item["presentacion"]
            hoja[f"E{r_cursor}"].alignment = Alignment(horizontal="center")
            hoja[f"E{r_cursor}"].font = Font(name="Segoe UI", size=9, color="475569")

            hoja[f"F{r_cursor}"].value = ""
            hoja[f"F{r_cursor}"].alignment = Alignment(horizontal="center")

            hoja[f"G{r_cursor}"].value = ""
            hoja[f"G{r_cursor}"].number_format = "#,##0.00"
            hoja[f"G{r_cursor}"].alignment = Alignment(horizontal="right")

            hoja[f"H{r_cursor}"].value = ""
            hoja[f"H{r_cursor}"].alignment = Alignment(horizontal="left")

            for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
                cell = hoja[f"{col}{r_cursor}"]
                cell.border = border_cell
                if row_fill:
                    cell.fill = row_fill

            hoja.row_dimensions[r_cursor].height = 20
            r_cursor += 1

        # Fila de Subtotal
        hoja.merge_cells(f"A{r_cursor}:C{r_cursor}")
        hoja[f"A{r_cursor}"].value = f"TOTAL {titulo.upper()}:"
        hoja[f"A{r_cursor}"].font = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
        hoja[f"A{r_cursor}"].alignment = Alignment(horizontal="right")

        hoja[f"D{r_cursor}"].value = f"=SUM(D{th_r + 1}:D{r_cursor - 1})"
        hoja[f"D{r_cursor}"].number_format = "#,##0.00"
        hoja[f"D{r_cursor}"].font = Font(name="Segoe UI", size=10, bold=True, color="006156")
        hoja[f"D{r_cursor}"].alignment = Alignment(horizontal="right")

        for col in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            c = hoja[f"{col}{r_cursor}"]
            c.border = Border(top=Side(border_style="thin", color="0F172A"), bottom=Side(border_style="medium", color="0F172A"))
            c.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        hoja.row_dimensions[r_cursor].height = 22
        return r_cursor + 2

    current_row = escribir_tabla_canal(ws, current_row, "Plaza de Mercado", "Perecederos, Carnes, Frutas y Verduras frescas", "B45309", lineas_mercado)
    current_row = escribir_tabla_canal(ws, current_row, "Supermercado y Abarrotes", "Secos, Lácteos industriales, Granos y Limpieza", "006156", lineas_super)
    current_row = escribir_tabla_canal(ws, current_row, "Proveedores Directos", "Distribuidoras, Panadería, Kéfir y Especiales", "4338CA", lineas_prov)
    if lineas_otros:
        current_row = escribir_tabla_canal(ws, current_row, "Otros Insumos y Servicios", "Descartables, envases y consumos varios", "475569", lineas_otros)

    # Firmas institucionales
    current_row += 1
    ws.merge_cells(f"A{current_row}:C{current_row}")
    ws[f"A{current_row}"].value = "____________________________________"
    ws[f"A{current_row}"].alignment = Alignment(horizontal="center")

    ws.merge_cells(f"E{current_row}:G{current_row}")
    ws[f"E{current_row}"].value = "____________________________________"
    ws[f"E{current_row}"].alignment = Alignment(horizontal="center")
    current_row += 1

    ws.merge_cells(f"A{current_row}:C{current_row}")
    ws[f"A{current_row}"].value = "Firma Solicitante Cocina"
    ws[f"A{current_row}"].font = Font(name="Segoe UI", size=9, bold=True, color="475569")
    ws[f"A{current_row}"].alignment = Alignment(horizontal="center")

    ws.merge_cells(f"E{current_row}:G{current_row}")
    ws[f"E{current_row}"].value = "Firma Gobernanta / Control de Suministros"
    ws[f"E{current_row}"].font = Font(name="Segoe UI", size=9, bold=True, color="475569")
    ws[f"E{current_row}"].alignment = Alignment(horizontal="center")

    # Guardar en memoria
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@app.get("/pedidos/{id_publico}/reporte/excel", tags=["Pedidos"])
@app.get("/pedidos/{id_publico}/reporte-excel", tags=["Pedidos"])
def ver_reporte_pedido_excel(id_publico: str, db: Session = Depends(get_db)):
    """
    Genera y descarga un archivo Excel (.xlsx) oficial de alto nivel con
    categorización por canales de compra, formato numérico profesional y campos de control.
    """
    pedido = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.id_publico == id_publico).first()

    if not pedido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pedido no encontrado"
        )

    buf = generar_excel_pedido(pedido)
    fecha_slug = pedido.fecha_solicitud.strftime("%Y%m%d_%H%M") if pedido.fecha_solicitud else "pedido"
    filename = f"Orden_Compra_{pedido.id_publico[:8].upper()}_{fecha_slug}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        }
    )


@app.get("/api/pedidos/pendientes/excel", tags=["Pedidos"])
def ver_consolidado_pendientes_excel(db: Session = Depends(get_db)):
    """
    Genera un archivo Excel (.xlsx) con el consolidado maestro de TODOS los pedidos pendientes,
    agrupando y sumando las cantidades de cada insumo por canal de compra para la gobernanta.
    """
    pedidos = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.estado == "pendiente").all()

    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado Pendientes"
    ws.views.sheetView[0].showGridLines = True

    TEAL_PRIMARY = "006156"
    BORDER_GRAY = "CBD5E1"
    thin_border = Side(border_style="thin", color=BORDER_GRAY)
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    ws.merge_cells("A1:G1")
    ws["A1"] = "CLÍNICA MONTALVO — DULCE ESPERA"
    ws["A1"].font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color=TEAL_PRIMARY, end_color=TEAL_PRIMARY, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:G2")
    ws["A2"] = f"CONSOLIDADO MAESTRO DE COMPRAS PENDIENTES — {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(name="Segoe UI", size=9, bold=True, color="E2E8F0")
    ws["A2"].fill = PatternFill(start_color=TEAL_PRIMARY, end_color=TEAL_PRIMARY, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Consolidar insumos por canal acumulando cantidades
    canales_acum = {
        "Mercado": {},
        "Super Mercado": {},
        "Proveedor": {},
        "Otros": {}
    }
    for ped in pedidos:
        for lin in ped.lineas:
            c = resolver_canal_compra(lin)
            if c not in canales_acum:
                c = "Otros"
            ins_id = lin.insumo_id_publico or "sin_id"
            ins_nombre = lin.insumo.nombre if lin.insumo else "Insumo"
            ins_pres = lin.insumo.presentacion if lin.insumo else "Unidad"
            ins_cat = lin.insumo.categoria if lin.insumo else "Otros"
            cant = float(lin.cantidad) if lin.cantidad else 0.0

            if ins_id not in canales_acum[c]:
                canales_acum[c][ins_id] = {
                    "nombre": ins_nombre,
                    "categoria": ins_cat,
                    "presentacion": ins_pres,
                    "cantidad": 0.0,
                    "pedidos_count": 0
                }
            canales_acum[c][ins_id]["cantidad"] += cant
            canales_acum[c][ins_id]["pedidos_count"] += 1

    col_widths = {"A": 6, "B": 38, "C": 20, "D": 15, "E": 14, "F": 16, "G": 28}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    current_r = 4
    for canal_name, color, sub in [
        ("Mercado", "B45309", "Perecederos, Carnes, Frutas y Verduras"),
        ("Super Mercado", "006156", "Secos, Granos, Lácteos y Abarrotes"),
        ("Proveedor", "4338CA", "Distribuidoras Especiales y Panadería"),
        ("Otros", "475569", "Consumos Varios y Descartables")
    ]:
        items_dict = canales_acum.get(canal_name, {})
        ws.merge_cells(f"A{current_r}:G{current_r}")
        ws[f"A{current_r}"].value = f"CANAL: {canal_name.upper()} — {sub}"
        ws[f"A{current_r}"].font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        ws[f"A{current_r}"].fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        ws[f"A{current_r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_r].height = 24
        current_r += 1

        th_r = current_r
        headers = [
            ("A", "N°", Alignment(horizontal="center")),
            ("B", "Descripción Insumo", Alignment(horizontal="left")),
            ("C", "Categoría", Alignment(horizontal="left")),
            ("D", "Cant. Total", Alignment(horizontal="right")),
            ("E", "Unidad", Alignment(horizontal="center")),
            ("F", "Solicitudes", Alignment(horizontal="center")),
            ("G", "Observaciones / Check", Alignment(horizontal="left"))
        ]
        for col, title, align in headers:
            c = ws[f"{col}{th_r}"]
            c.value = title
            c.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            c.alignment = align
            c.border = border_cell
        ws.row_dimensions[th_r].height = 20
        current_r += 1

        if not items_dict:
            ws.merge_cells(f"A{current_r}:G{current_r}")
            ws[f"A{current_r}"].value = f"No hay insumos pendientes para {canal_name}."
            ws[f"A{current_r}"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
            ws[f"A{current_r}"].alignment = Alignment(horizontal="center")
            for col in ["A", "B", "C", "D", "E", "F", "G"]:
                ws[f"{col}{current_r}"].border = border_cell
            current_r += 2
            continue

        items_sorted = sorted(items_dict.values(), key=lambda x: (x["categoria"], x["nombre"]))
        for idx, it in enumerate(items_sorted, 1):
            is_zebra = (idx % 2 == 0)
            row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if is_zebra else None

            ws[f"A{current_r}"].value = idx
            ws[f"A{current_r}"].alignment = Alignment(horizontal="center")
            ws[f"A{current_r}"].font = Font(name="Segoe UI", size=9, bold=True, color="64748B")

            ws[f"B{current_r}"].value = it["nombre"]
            ws[f"B{current_r}"].font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

            ws[f"C{current_r}"].value = it["categoria"]
            ws[f"C{current_r}"].font = Font(name="Segoe UI", size=9, color="475569")

            ws[f"D{current_r}"].value = it["cantidad"]
            ws[f"D{current_r}"].number_format = "#,##0.00"
            ws[f"D{current_r}"].alignment = Alignment(horizontal="right")
            ws[f"D{current_r}"].font = Font(name="Segoe UI", size=10, bold=True, color="006156")

            ws[f"E{current_r}"].value = it["presentacion"]
            ws[f"E{current_r}"].alignment = Alignment(horizontal="center")
            ws[f"E{current_r}"].font = Font(name="Segoe UI", size=9, color="475569")

            ws[f"F{current_r}"].value = f"{it['pedidos_count']} ped."
            ws[f"F{current_r}"].alignment = Alignment(horizontal="center")
            ws[f"F{current_r}"].font = Font(name="Segoe UI", size=9, color="64748B")

            ws[f"G{current_r}"].value = ""

            for col in ["A", "B", "C", "D", "E", "F", "G"]:
                c = ws[f"{col}{current_r}"]
                c.border = border_cell
                if row_fill:
                    c.fill = row_fill
            ws.row_dimensions[current_r].height = 20
            current_r += 1

        # Subtotal
        ws.merge_cells(f"A{current_r}:C{current_r}")
        ws[f"A{current_r}"].value = f"TOTAL {canal_name.upper()}:"
        ws[f"A{current_r}"].font = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
        ws[f"A{current_r}"].alignment = Alignment(horizontal="right")

        ws[f"D{current_r}"].value = f"=SUM(D{th_r + 1}:D{current_r - 1})"
        ws[f"D{current_r}"].number_format = "#,##0.00"
        ws[f"D{current_r}"].font = Font(name="Segoe UI", size=10, bold=True, color="006156")
        ws[f"D{current_r}"].alignment = Alignment(horizontal="right")

        for col in ["A", "B", "C", "D", "E", "F", "G"]:
            c = ws[f"{col}{current_r}"]
            c.border = Border(top=Side(border_style="thin", color="0F172A"), bottom=Side(border_style="medium", color="0F172A"))
            c.fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        ws.row_dimensions[current_r].height = 22
        current_r += 2

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"Consolidado_Pendientes_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        }
    )


@app.get("/pedidos/{id_publico}/reporte-abastecimiento", response_class=HTMLResponse, tags=["Pedidos"])
def ver_reporte_pedido_abastecimiento(id_publico: str, db: Session = Depends(get_db)):
    """
    Genera una lista de verificación y control de calidad (HACCP) para la recepción
    de insumos del pedido en la cocina.
    """
    pedido = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.id_publico == id_publico).first()

    if not pedido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pedido no encontrado"
        )

    fecha_str = pedido.fecha_solicitud.strftime("%Y-%m-%d %H:%M") if pedido.fecha_solicitud else "N/A"

    rows_html = ""
    for idx, linea in enumerate(pedido.lineas, 1):
        nombre = esc(linea.insumo.nombre) if linea.insumo else "Insumo sin nombre"
        presentacion = esc(linea.insumo.presentacion) if linea.insumo else "Unidades"
        categoria = esc(linea.insumo.categoria) if linea.insumo else "Otros"
        cantidad_val = float(linea.cantidad) if linea.cantidad else 0.0
        id_insumo = linea.insumo_id_publico or "N/A"

        rows_html += f"""
        <tr style="border-bottom: 1px solid #cbd5e1;">
            <td style="padding: 12px 8px; text-align: center; font-weight: bold; color: #64748b; font-size: 13px;">{idx}</td>
            <td style="padding: 12px 8px; font-weight: 700; color: #0f172a; font-size: 13px;">{nombre}</td>
            <td style="padding: 12px 8px; text-align: right; font-weight: 800; color: #006156; font-size: 14px;">{cantidad_val:.2f}</td>
            <td style="padding: 12px 8px; color: #475569; font-size: 12px;">{presentacion}</td>
            <td style="padding: 12px 8px; text-align: center;"><span style="display: inline-block; width: 14px; height: 14px; border: 1.5px solid #94a3b8; border-radius: 2px;"></span></td>
            <td style="padding: 12px 8px; border-bottom: 1px solid #cbd5e1; width: 120px;"></td>
            <td style="padding: 12px 8px; border-bottom: 1px solid #cbd5e1; width: 100px;"></td>
        </tr>
        """

    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Lista de Verificación de Calidad - Dulce Espera</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: #0f172a;
            background-color: #f1f5f9;
            margin: 0;
            padding: 30px 20px;
        }}
        .container {{
            max-width: 950px;
            margin: 0 auto;
            background-color: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 20px;
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.05), 0 4px 6px -2px rgba(0, 0, 0, 0.02);
            padding: 45px;
            box-sizing: border-box;
        }}
        .header {{
            display: flex;
            justify-content: space-between;
            align-items: flex-start;
            border-bottom: 2px solid #cbd5e1;
            padding-bottom: 24px;
            margin-bottom: 30px;
        }}
        .brand {{
            display: flex;
            align-items: center;
            gap: 16px;
        }}
        .brand-title {{
            font-size: 26px;
            font-weight: 900;
            color: #006156;
            margin: 0;
            letter-spacing: -0.03em;
        }}
        .doc-type {{
            color: #475569;
            font-size: 12px;
            font-weight: 700;
            display: inline-block;
            margin-top: 6px;
            letter-spacing: 0.05em;
        }}
        .meta {{
            text-align: right;
            font-size: 13px;
            color: #475569;
            line-height: 1.6;
        }}
        .meta strong {{
            color: #0f172a;
        }}
        .details-table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 5px;
        }}
        .details-table th {{
            background-color: #f8fafc;
            color: #475569;
            font-weight: 800;
            text-transform: uppercase;
            font-size: 10px;
            letter-spacing: 0.05em;
            padding: 12px 8px;
            border-bottom: 2px solid #cbd5e1;
        }}
        .print-btn {{
            display: inline-flex;
            align-items: center;
            justify-content: center;
            background-color: #006156;
            color: #ffffff;
            border: none;
            border-radius: 10px;
            padding: 12px 24px;
            font-size: 14px;
            font-weight: 700;
            cursor: pointer;
            transition: all 0.2s ease;
            margin-bottom: 20px;
            gap: 8px;
            text-decoration: none;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
        }}
        .print-btn:hover {{
            background-color: #004d45;
            transform: translateY(-1px);
        }}
        .no-print {{
            display: flex;
            justify-content: flex-end;
            max-width: 950px;
            margin: 0 auto;
        }}
        .footer {{
            border-top: 1px solid #e2e8f0;
            padding-top: 20px;
            text-align: center;
            font-size: 12px;
            color: #64748b;
            margin-top: 50px;
            line-height: 1.5;
        }}
        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .container {{
                border: none;
                box-shadow: none;
                padding: 0;
                max-width: 100%;
            }}
            .no-print {{
                display: none !important;
            }}
        }}
    </style>
</head>
<body>
    <div class="no-print">
        <button class="print-btn" onclick="window.print()">
            <svg style="width: 18px; height: 18px;" fill="none" stroke="currentColor" viewBox="0 0 24 24" xmlns="http://www.w3.org/2000/svg">
                <path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M17 17h2a2 2 0 002-2v-4a2 2 0 00-2-2H5a2 2 0 00-2 2v4a2 2 0 002 2h2m2 4h10a2 2 0 002-2v-4a2 2 0 00-2-2H9a2 2 0 00-2 2v4a2 2 0 002 2zm8-12V5a2 2 0 00-2-2H9a2 2 0 00-2 2v4h10z"></path>
            </svg>
            Imprimir Formato
        </button>
    </div>
    <div class="container">
        <div class="header">
            <div class="brand">
                <img src="https://dulce-espera-inventario.vercel.app/logo.svg" alt="Logo" style="width: 75px; height: 75px; object-fit: contain;" onerror="this.style.display='none'">
                <div>
                    <h1 class="brand-title">DULCE ESPERA</h1>
                    <span class="doc-type">LISTA DE VERIFICACIÓN Y CONTROL DE CALIDAD DE RECEPCIÓN (HACCP)</span>
                </div>
            </div>
            <div class="meta">
                <div><strong>ID Pedido:</strong> {pedido.id_publico[:8].upper()}</div>
                <div style="font-size: 10px; color: #94a3b8; margin-bottom: 2px;">UUID: {pedido.id_publico}</div>
                <div><strong>Fecha Pedido:</strong> {fecha_str}</div>
                <div><strong>Solicitante:</strong> {esc(pedido.solicitante)}</div>
                <div><strong>Estado Actual:</strong> <span style="text-transform: uppercase; font-weight: 800; color: #006156;">{esc(pedido.estado)}</span></div>
            </div>
        </div>

        <div style="padding: 12px 15px; border-left: 3px solid #006156; background-color: #f0faf9; color: #006156; font-size: 12px; font-weight: 500; margin-bottom: 25px; line-height: 1.5;">
            <strong>Instrucciones para el Operario de Cocina:</strong> Verifique el estado físico, empaque, limpieza y temperatura de cada insumo al recibirlo. Marque la conformidad y anote el lote/fecha de vencimiento correspondiente para asegurar la trazabilidad.
        </div>

        <table class="details-table">
            <thead>
                <tr>
                    <th style="width: 45px; text-align: center;">Item</th>
                    <th style="text-align: left;">Descripción del Insumo</th>
                    <th style="text-align: right; width: 90px;">Cant.</th>
                    <th style="text-align: left; width: 90px;">Unidad</th>
                    <th style="text-align: center; width: 80px;">Conforme</th>
                    <th style="text-align: left; width: 120px;">Lote / Temp.</th>
                    <th style="text-align: left; width: 100px;">Fec. Vencimiento</th>
                </tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>

        <div style="margin-top: 60px; display: flex; justify-content: space-between; gap: 30px;">
            <div style="text-align: center; flex: 1; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 11px; color: #475569; font-weight: 600;">
                Firma Responsable de Control de Calidad
            </div>
            <div style="text-align: center; flex: 1; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 11px; color: #475569; font-weight: 600;">
                Firma Encargado de Recepción de Cocina
            </div>
        </div>

        <div class="footer">
            Este documento físico de control de calidad debe archivarse para auditorías sanitarias y control interno de la cocina.<br>
            &copy; {datetime.now().year} Dulce Espera.
        </div>
    </div>
</body>
</html>"""
    return HTMLResponse(content=html_content)


@app.get("/pedidos/{id_publico}/reporte-categorias", response_class=HTMLResponse, tags=["Pedidos"])
def ver_reporte_pedido_categorias(id_publico: str, db: Session = Depends(get_db)):
    """
    Genera un informe administrativo del pedido clasificado y agrupado de forma
    granular por las categorías de insumos de la base de datos.
    """
    pedido = db.query(models.Pedido).options(
        joinedload(models.Pedido.lineas).joinedload(models.DetallePedido.insumo)
    ).filter(models.Pedido.id_publico == id_publico).first()

    if not pedido:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Pedido no encontrado"
        )

    fecha_str = pedido.fecha_solicitud.strftime("%Y-%m-%d %H:%M") if pedido.fecha_solicitud else "N/A"

    # Group lines by category
    grouped_items = {}
    for linea in pedido.lineas:
        nombre = esc(linea.insumo.nombre) if linea.insumo else "Insumo sin nombre"
        presentacion = esc(linea.insumo.presentacion) if linea.insumo else "Unidades"
        categoria = esc(linea.insumo.categoria) if (linea.insumo and linea.insumo.categoria) else "Sin Categoría"
        cantidad_val = float(linea.cantidad) if linea.cantidad else 0.0
        id_insumo = linea.insumo_id_publico or "N/A"

        item_data = {
            "nombre": nombre,
            "presentacion": presentacion,
            "cantidad": cantidad_val,
            "id_insumo": id_insumo
        }

        if categoria not in grouped_items:
            grouped_items[categoria] = []
        grouped_items[categoria].append(item_data)

    # Render HTML for each category
    html_categories = ""
    for cat_name in sorted(grouped_items.keys()):
        items_list = grouped_items[cat_name]
        html_categories += f"""
        <div style="margin-bottom: 25px;">
            <h3 style="font-size: 13px; margin-bottom: 10px; font-weight: 800; color: #006156; border-left: 3px solid #39ada3; padding-left: 8px; text-transform: uppercase; letter-spacing: 0.05em;">
                Categoría: {cat_name} ({len(items_list)} insumo{'' if len(items_list) == 1 else 's'})
            </h3>
            <table class="details-table">
                <thead>
                    <tr>
                        <th style="width: 45px; text-align: center;">Item</th>
                        <th style="width: 80px; text-align: left;">Código</th>
                        <th style="text-align: left;">Descripción Insumo</th>
                        <th style="text-align: right; width: 120px;">Cant. Solicitada</th>
                        <th style="text-align: left; width: 120px;">Unidad</th>
                    </tr>
                </thead>
                <tbody>
        """
        for idx, item in enumerate(items_list, 1):
            html_categories += f"""
                    <tr style="border-bottom: 1px solid #e2e8f0;">
                        <td style="padding: 9px 10px; text-align: center; font-weight: bold; color: #64748b; font-size: 12px;">{idx}</td>
                        <td style="padding: 9px 10px; color: #64748b; font-size: 11px; font-family: monospace;">{item['id_insumo'][:8]}</td>
                        <td style="padding: 9px 10px; font-weight: 700; color: #0f172a; text-align: left; font-size: 12px;">{item['nombre']}</td>
                        <td style="padding: 9px 10px; text-align: right; font-weight: 800; color: #006156; font-size: 13px;">{item['cantidad']:.2f}</td>
                        <td style="padding: 9px 10px; color: #475569; font-weight: 500; font-size: 12px; text-align: left;">{item['presentacion']}</td>
                    </tr>
            """
        html_categories += """
                </tbody>
            </table>
        </div>
        """

    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Reporte por Categorías - Dulce Espera</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: #0f172a;
            background-color: #f1f5f9;
            margin: 0;
            padding: 30px 20px;
        }}
        .container {{
            max-width: 900px;
            margin: 0 auto;
            background-color: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 20px;
            box-shadow: 0 10px 15px -3px rgba(0, 0, 0, 0.05), 0 4px 6px -2px rgba(0, 0, 0, 0.02);
            padding: 45px;
            box-sizing: border-box;
        }}
        .header {{
            display: flex;
            justify-content: space-between;
            align-items: flex-start;
            border-bottom: 2px solid #cbd5e1;
            padding-bottom: 24px;
            margin-bottom: 30px;
        }}
        .brand {{
            display: flex;
            align-items: center;
            gap: 16px;
        }}
        .brand-title {{
            font-size: 26px;
            font-weight: 900;
            color: #006156;
            margin: 0;
            letter-spacing: -0.03em;
        }}
        .doc-type {{
            color: #475569;
            font-size: 12px;
            font-weight: 700;
            display: inline-block;
            margin-top: 6px;
            letter-spacing: 0.05em;
        }}
        .meta {{
            text-align: right;
            font-size: 13px;
            color: #475569;
            line-height: 1.6;
        }}
        .meta strong {{
            color: #0f172a;
        }}
        .details-table {{
            width: 100%;
            border-collapse: collapse;
            margin-top: 5px;
        }}
        .details-table th {{
            background-color: #f8fafc;
            color: #475569;
            font-weight: 800;
            text-transform: uppercase;
            font-size: 10px;
            letter-spacing: 0.05em;
            padding: 10px;
            border-bottom: 2px solid #cbd5e1;
        }}
        .print-btn {{
            display: inline-flex;
            align-items: center;
            justify-content: center;
            background-color: #006156;
            color: #ffffff;
            border: none;
            border-radius: 10px;
            padding: 12px 24px;
            font-size: 14px;
            font-weight: 700;
            cursor: pointer;
            transition: all 0.2s ease;
            margin-bottom: 20px;
            gap: 8px;
            text-decoration: none;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
        }}
        .print-btn:hover {{
            background-color: #004d45;
            transform: translateY(-1px);
        }}
        .no-print {{
            display: flex;
            justify-content: flex-end;
            max-width: 900px;
            margin: 0 auto;
        }}
        .footer {{
            border-top: 1px solid #e2e8f0;
            padding-top: 20px;
            text-align: center;
            font-size: 12px;
            color: #64748b;
            margin-top: 50px;
            line-height: 1.5;
        }}
        @media print {{
            body {{
                background-color: #ffffff;
                padding: 0;
            }}
            .container {{
                border: none;
                box-shadow: none;
                padding: 0;
                max-width: 100%;
            }}
            .no-print {{
                display: none !important;
            }}
        }}
    </style>
</head>
<body>
    <div class="no-print">
        <button class="print-btn" onclick="window.print()">
            <svg style="width: 18px; height: 18px;" fill="none" stroke="currentColor" viewBox="0 0 24 24" xmlns="http://www.w3.org/2000/svg">
                <path stroke-linecap="round" stroke-linejoin="round" stroke-width="2" d="M17 17h2a2 2 0 002-2v-4a2 2 0 00-2-2H5a2 2 0 00-2 2v4a2 2 0 002 2h2m2 4h10a2 2 0 002-2v-4a2 2 0 00-2-2H9a2 2 0 00-2 2v4a2 2 0 002 2zm8-12V5a2 2 0 00-2-2H9a2 2 0 00-2 2v4h10z"></path>
            </svg>
            Imprimir Reporte
        </button>
    </div>
    <div class="container">
        <div class="header">
            <div class="brand">
                <img src="https://dulce-espera-inventario.vercel.app/logo.svg" alt="Logo" style="width: 75px; height: 75px; object-fit: contain;" onerror="this.style.display='none'">
                <div>
                    <h1 class="brand-title">DULCE ESPERA</h1>
                    <span class="doc-type">INFORME DE COMPRAS CLASIFICADO POR CATEGORÍAS</span>
                </div>
            </div>
            <div class="meta">
                <div><strong>ID Pedido:</strong> {pedido.id_publico[:8].upper()}</div>
                <div style="font-size: 10px; color: #94a3b8; margin-bottom: 2px;">UUID: {pedido.id_publico}</div>
                <div><strong>Fecha Pedido:</strong> {fecha_str}</div>
                <div><strong>Solicitante:</strong> {esc(pedido.solicitante)}</div>
                <div><strong>Estado Actual:</strong> <span style="text-transform: uppercase; font-weight: 800; color: #006156;">{esc(pedido.estado)}</span></div>
            </div>
        </div>

        <div style="display: flex; gap: 40px; margin-bottom: 30px; border-bottom: 2px solid #cbd5e1; padding-bottom: 12px; flex-wrap: wrap;">
            <div style="font-size: 12px; color: #475569; font-weight: 600;">
                TOTAL INSUMOS PEDIDO: <span style="font-size: 15px; font-weight: 800; color: #0f172a; margin-left: 4px;">{len(pedido.lineas)}</span>
            </div>
            <div style="font-size: 12px; color: #475569; font-weight: 600;">
                CATEGORÍAS PRESENTES: <span style="font-size: 15px; font-weight: 800; color: #006156; margin-left: 4px;">{len(grouped_items)}</span>
            </div>
        </div>

        {html_categories}

        <div style="margin-top: 60px; display: flex; justify-content: space-between; gap: 30px;">
            <div style="text-align: center; flex: 1; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 11px; color: #475569; font-weight: 600;">
                Firma Solicitante Cocina
            </div>
            <div style="text-align: center; flex: 1; border-top: 1px solid #cbd5e1; padding-top: 8px; font-size: 11px; color: #475569; font-weight: 600;">
                Autorizado Administración / Compras
            </div>
        </div>

        <div class="footer">
            Este informe desglosa de manera detallada las compras solicitadas de acuerdo con la categorización del catálogo.<br>
            &copy; {datetime.now().year} Dulce Espera.
        </div>
    </div>
</body>
</html>"""
    return HTMLResponse(content=html_content)

